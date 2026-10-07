"""모델 → xlsx (tasks/0008 4.3·4.6). 모델을 파일로 옮길 뿐이다 — 무엇을 싣는가는 모델이 정한다.

- 날짜와 시각은 글자 그대로 (엑셀이 날짜로 바꿔 읽지 않게). 수는 수로. 합계도 값으로 (수식을 쓰지 않는다).
- "=" 로 시작하는 글자는 수식이 되지 않게 글자로 적는다. 엑셀이 받지 않는 제어 문자는 뺀다.
- 표시는 바탕색과 글자 표시를 같이 (모델의 칸에 글자 표시가 이미 들어 있다). 머리글 행은 틀 고정, 긴 표에는 자동 필터.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FILL = {"pending": "FFF2CC", "illegible": "F8CBAD", "printed": "EDEDED", "meta": "DDEBF7", "no_doc": "D9D9D9",
        "unknown": "E2EFDA", "mismatch": "F4B6B6", "h": "D9E1F2"}
_FILLS = {k: PatternFill(fill_type="solid", start_color=v, end_color=v) for k, v in FILL.items()}
_BOLD, _ITALIC, _TITLE = Font(bold=True), Font(italic=True, color="595959"), Font(bold=True, size=12)
_MISMATCH = Font(bold=True, color="9C0006")
MAX_WIDTH = 48


def text(v):
    """엑셀에 적을 값: 글자에서 제어 문자를 뺀다 (모델의 값은 그대로 두고 쓸 때만)."""
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def write_book(book: dict, path, made_at: str, version: str) -> None:
    """책 하나를 path(경로 또는 바이트를 받는 파일 객체 — 내려받기)에 쓴다 (원자적으로 바꾸는 것은 부른 쪽 — export/writer.py)."""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in book["sheets"]:
        ws = wb.create_sheet(sheet["name"])
        widths: dict[int, int] = {}
        for r, row in enumerate(sheet["rows"], start=1):
            for c, (value, style) in enumerate(row, start=1):
                tokens = style.split() if style else []
                if "stamp_time" in tokens:
                    value = made_at
                elif "stamp_version" in tokens:
                    value = version
                if value is None and not tokens:
                    continue
                v = text(value)
                cl = ws.cell(row=r, column=c, value=v)
                if isinstance(v, str) and cl.data_type in ("f", "e"):
                    cl.data_type = "s"                       # 수식·오류 값이 아니다 (적힌 글자 그대로 — "=…", "#N/A")
                for t in tokens:
                    if t in _FILLS:
                        cl.fill = _FILLS[t]
                if "h" in tokens:
                    cl.font = _BOLD
                elif "title" in tokens:
                    cl.font = _TITLE
                elif "note" in tokens:
                    cl.font = _ITALIC
                if "mismatch" in tokens:
                    cl.font = _MISMATCH
                if value is not None:
                    widths[c] = max(widths.get(c, 0), min(MAX_WIDTH, len(str(value)) + 2))
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = max(6, w)
        if sheet.get("freeze"):
            ws.freeze_panes = f"A{sheet['freeze'] + 1}"
        if sheet.get("filter") and sheet["rows"]:
            ws.auto_filter.ref = f"A1:{get_column_letter(max(len(r) for r in sheet['rows']))}{len(sheet['rows'])}"
    wb.save(path if hasattr(path, "write") else str(path))


def read_values(path: Path) -> dict[str, list[list]]:
    """되읽기 (시험·확인용): 시트 이름 → 행들의 값 (빈 칸은 None, 행 끝의 빈 칸은 뺀다)."""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            vals = list(row)
            while vals and vals[-1] is None:
                vals.pop()
            rows.append(vals)
        while rows and not rows[-1]:
            rows.pop()
        out[ws.title] = rows
    wb.close()
    return out
