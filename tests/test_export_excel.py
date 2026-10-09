"""엑셀 내보내기 (tasks/0008 단계 2): 내보내기 모델(DB → 시트들)과 일별 파일.

모델이 맞는가(DB·정답과)와 파일이 모델과 같은가(되읽어서)를 따로 본다 — 파일의 바이트는 견주지 않는다 (xlsx 안에 만든 시각이 있다).
세션 픽스처(기본 묶음의 null·oracle, 가동 일보 묶음의 oracle)를 같이 쓰고, 새로 돌리는 것은 하루치 셋뿐이다:
가짜 인식기(신뢰도가 기준 아래인 값), 가동 일보의 null, --display-names.
"""
from __future__ import annotations

import sqlite3
from collections import Counter
from pathlib import Path

import pytest
import yaml

from conftest import clone_db
from minedocscan.config import Settings
from minedocscan.export import labels as L
from minedocscan.export.daily import daily_book
from minedocscan.export.model import UNSURE, field_cell, typed
from minedocscan.export.writer import (
    RECORD_NAME,
    ExportError,
    check_out_dir,
    daily_path,
    export_excel,
    format_result,
    monthly_path,
)
from minedocscan.export.xlsx import read_values
from minedocscan.forms.equipment import is_equipment_row, layout
from minedocscan.forms.sitepack import SitePack
from minedocscan.forms.template import Template, display_of
from minedocscan.pipeline import Pipeline
from minedocscan.recognize import Recognition, load_answers_json
from minedocscan.recognize.builtin import OracleRecognizer
from minedocscan.review.store import review_from_field, save
from minedocscan.store.db import read_txn
from minedocscan.tools.synth import generate


def daily(paths) -> list[str]:
    """일별 파일만 (월별 파일은 단계 3 — test_export_monthly.py)."""
    return sorted(p for p in paths if p.startswith("daily/"))


def book(con, site, day, **kw) -> dict:
    with read_txn(con):
        return daily_book(con, site, day, **kw)


def sheet(b: dict, name: str) -> dict:
    return next(s for s in b["sheets"] if s["name"] == name)


def blocks(s: dict) -> dict[str, list[list]]:
    """양식 시트 → 쪽 ID → 그 쪽의 블록 (머리 행부터 다음 쪽의 머리 전까지)."""
    out, cur = {}, None
    for row in s["rows"]:
        if row and row[0] == [L.PAGE_HEAD, "h"]:
            cur = out.setdefault(row[2][0], [])
        if cur is not None:
            cur.append(row)
    return out


def days_of(con) -> list[str]:
    return [r[0] for r in con.execute("SELECT DISTINCT work_date FROM doc_page WHERE work_date IS NOT NULL ORDER BY 1")]


def form_cells(site, con, b: dict):
    """양식 시트의 칸마다 (쪽 ID, 표, 행 번호, 열 이름, 칸) — 템플릿의 행·열 순서대로 모델을 읽는다 (순서가 종이와 같다는 것도 본다)."""
    pages = {r["page_id"]: r for r in con.execute("SELECT * FROM doc_page WHERE status = 'loaded'")}
    for s in b["sheets"]:
        for pid, rows in blocks(s).items():
            tpl = site.templates[pages[pid]["template_name"]]
            i = 1
            while len(rows[i]) == 2 and rows[i][1][1] == "note":        # 핸들러가 말한 줄
                i += 1
            if tpl.fields:
                assert rows[i][0][0] == L.FIELDS_HEAD
                for k, f in enumerate(tpl.fields):
                    row = rows[i + 1 + k]
                    assert row[0][0] == display_of(f, f["name"])
                    yield pid, "fields", -1, f["name"], row[1]
                i += 1 + len(tpl.fields)
            for reg in tpl.regions:
                assert rows[i][0][0] == display_of(reg, reg["name"])
                cols = sorted(reg["columns"], key=lambda c: c["idx"])
                assert [c[0] for c in rows[i + 1]] == [L.ROW_HEAD, *(display_of(c, c["name"]) for c in cols)]
                for k, r in enumerate(sorted(reg["rows"], key=lambda r: r["row"])):
                    row = rows[i + 2 + k]
                    assert row[0][0] == (display_of(r, str(r.get("key", "") or "")) or None)     # 여백 행은 이름이 없다
                    for j, c in enumerate(cols):
                        yield pid, reg["name"], r["row"], c["name"], row[1 + j]
                i += 2 + len(reg["rows"])


def check_form_cells(site, con, b: dict) -> Counter:
    """4.2 의 첫째 표대로인가 — 칸 하나하나를 그 doc_field 행·쪽 메타와 견준다. 돌려주는 값: 상태별 칸 수."""
    meta = {(r["page_id"], r["meta_key"]): r["value"] for r in con.execute("SELECT * FROM doc_page_meta WHERE value IS NOT NULL")}
    pages = {r["page_id"]: r["template_name"] for r in con.execute("SELECT page_id, template_name FROM doc_page")}
    seen = Counter()
    for pid, region, row, name, (value, style) in form_cells(site, con, b):
        f = dict(con.execute("SELECT * FROM doc_field WHERE field_id = ?", (f"{pid}:{region}:{name}:{row}",)).fetchone())
        tpl = site.templates[pages[pid]]
        mk = next((x.get("meta_key") for x in tpl.fields if x["name"] == name), None) if region == "fields" else None
        seen[style] += 1
        if f["kind"] == "printed":
            assert (value, style) == (f["value_final"] or None, "printed")
        elif mk and meta.get((pid, mk)) is not None:
            assert (value, style) == (meta[(pid, mk)], "meta")
        elif style in ("value", "mark"):                         # 확정된 값 — 기계가 자동 적재했거나 사람이 검수했다
            assert f["review_status"] != "pending" and f["has_value"] == 1
            assert value == (L.CHECK_MARK if f["kind"] == "checkmark" else typed(f["value_final"], f["format"]))
        elif style == "pending":
            assert f["review_status"] == "pending" and f["reviewed_by"] is None and value == L.PENDING_MARK
        elif style == "illegible":
            assert f["review_status"] == "pending" and f["reviewed_by"] and value == L.ILLEGIBLE_MARK
        elif style == "empty":                                   # 빈 칸, 또는 핸들러가 비운 ✓ 칸(점검하지 않은 쪽·여백 행)
            assert value is None
            assert f["has_value"] == 0 and f["review_status"] != "pending" or (
                f["kind"] == "checkmark" and tpl.handler == "inspection" and f["review_status"] == "pending")
        elif style == "present":                                 # 서명처럼 유무만, 또는 작업 표의 검수 대기 글자 칸
            assert value == L.PRESENT_MARK
            assert (f["has_value"] == 1 and f["review_status"] != "pending" and not f["value_final"]) or (
                tpl.handler == "usage" and f["kind"] == "handwritten_text" and f["review_status"] == "pending")
        else:
            raise AssertionError((pid, region, row, name, style))
    return seen


BUSINESS = {"haul_log": ("SELECT COUNT(*) FROM prod_haul WHERE work_date = ? AND source_role = 'log'"),
            "haul_matrix": ("SELECT COUNT(*) FROM prod_haul WHERE work_date = ? AND source_role = 'matrix'"),
            "xcheck_haul": "SELECT COUNT(*) FROM xcheck_haul WHERE work_date = ?",
            "assignment": "SELECT COUNT(*) FROM eq_assignment_obs WHERE work_date = ?",
            "usage": "SELECT COUNT(*) FROM eq_usage_daily WHERE work_date = ?",
            "tally": "SELECT COUNT(*) FROM prod_tally WHERE work_date = ?",
            "xcheck_usage": "SELECT COUNT(*) FROM xcheck_usage WHERE work_date = ?",
            "inspection": "SELECT COUNT(*) FROM insp_daily WHERE inspection_date = ?"}


def business_rows(b: dict, key: str) -> list[dict]:
    """업무 시트 → 행마다 {머리글: 값}. 시트가 없으면 []."""
    s = next((s for s in b["sheets"] if s["name"] == L.SHEETS[key]), None)
    if s is None:
        return []
    names = [c[0] for c in s["rows"][0]]
    return [{n: c[0] for n, c in zip(names, r, strict=True)} for r in s["rows"][1:]]


def state_of(con, fid: str | None) -> str:
    """시험 쪽에서 따로 세는 필드의 상태 (4.2): ID 가 없으면 견준 것이 없다(확정), ID 가 있는데 행이 없으면 모른다."""
    if fid is None:
        return "value"
    f = con.execute("SELECT review_status, reviewed_by, has_value FROM doc_field WHERE field_id = ?", (fid,)).fetchone()
    if f is None:
        return "pending"
    if f["review_status"] == "pending":
        return "illegible" if f["reviewed_by"] else "pending"
    return "value" if f["has_value"] else "empty"


def all_sure(states) -> bool:
    return all(s not in ("pending", "illegible") for s in states)


def check_business(con, b: dict, day: str, site=None) -> None:
    """업무 시트의 행 수 = 업무 테이블의 그 날짜 행 수, 값 열은 4.2 의 둘째 표대로 (확정일 때만) — 기대하는 상태는 doc_field 에서
    따로 센다 (모델의 상태 열을 믿지 않는다)."""
    from minedocscan.handlers import get_handler
    from minedocscan.handlers.usage import is_shift_cell
    from minedocscan.validate.usage import check_cells

    for key, sql in BUSINESS.items():
        assert len(business_rows(b, key)) == con.execute(sql, (day,)).fetchone()[0], key
    src = {f"{r['source_name']}#{r['page_no']}": r["page_id"] for r in con.execute(
        "SELECT p.page_id, p.page_no, d.source_name FROM doc_page p JOIN doc_document d ON p.document_id = d.document_id")}
    for role in ("log", "matrix"):
        for r in business_rows(b, f"haul_{role}"):
            h = con.execute("SELECT * FROM prod_haul WHERE haul_id = ?", (r[L.FIELD_ID],)).fetchone()
            st = state_of(con, h["source_field_id"])
            assert r[L.STATE] == L.ROW_STATE[st]
            assert r["횟수"] == (h["trips"] if st == "value" else None)
    for r in business_rows(b, "tally"):
        t = con.execute("SELECT * FROM prod_tally WHERE tally_id = ?", (r[L.FIELD_ID],)).fetchone()
        st = state_of(con, t["source_field_id"])
        assert r["수"] == (t["count"] if st == "value" else None) and r[L.STATE] == L.ROW_STATE[st]
    # 교차검증: 그 날짜·자리·광종·편의 운반 행이 전부 확정일 때만 횟수, 아니면 판정에 (잠정)
    for r in business_rows(b, "xcheck_haul"):
        if r["자리"] == L.UNRESOLVED_SLOT:                    # 자리 미정: 그 쪽(출처)의 일보 행이 전부 확정일 때만 (행이 없으면 확정 아님)
            assert r["차량번호"] is None and r["작성자"] is None
            pids = [src[x] for x in r[L.SOURCE].split(", ")]
            hs = [h[0] for h in con.execute(
                f"SELECT source_field_id FROM prod_haul WHERE source_role = 'log' AND material = ? AND level = ? AND page_id IN "
                f"({','.join('?' * len(pids))})", (r["광종"], r["편"], *pids))]
            ok = bool(hs) and all_sure(state_of(con, f) for f in hs)
            xs = con.execute("SELECT * FROM xcheck_haul WHERE work_date = ? AND slot LIKE 'unresolved:%' AND material = ? "
                             "AND level = ?", (day, r["광종"], r["편"])).fetchall()
            assert any((r["일보 횟수"], r["행렬 횟수"]) == ((x["log_trips"], x["matrix_trips"]) if ok else (None, None))
                       and r["판정"] == L.XCHECK_STATUS[x["status"]] + ("" if ok else L.PROVISIONAL) for x in xs), r
            assert (r[L.STATE] == L.ROW_STATE["value"]) if ok else r[L.STATE] in (L.ROW_STATE["pending"], L.ROW_STATE["illegible"])
            continue
        x = con.execute("SELECT * FROM xcheck_haul WHERE work_date = ? AND slot = ? AND material = ? AND level = ?",
                        (day, r["자리"], r["광종"], r["편"])).fetchone()
        hs = [h[0] for h in con.execute("SELECT source_field_id FROM prod_haul WHERE work_date = ? AND slot = ? AND material = ? "
                                        "AND level = ?", (day, r["자리"], r["광종"], r["편"]))]
        ok = all_sure(state_of(con, f) for f in hs)
        assert (r["일보 횟수"], r["행렬 횟수"]) == ((x["log_trips"], x["matrix_trips"]) if ok else (None, None))
        assert r["판정"] == L.XCHECK_STATUS[x["status"]] + ("" if ok else L.PROVISIONAL)
    for r in business_rows(b, "assignment"):
        x = con.execute("SELECT * FROM eq_assignment_obs WHERE work_date = ? AND slot = ?", (day, r["자리"])).fetchone()
        assert (r["차량번호"], r["작성자"], r["머리글 차량번호"], r["머리글 작성자"]) == (
            x["vehicle_no"], x["operator"], x["header_vehicle_no"], x["header_operator"])
    # 가동 기록: 값마다 그 칸, 가동 시간은 근거가 된 칸 전부
    for r in business_rows(b, "usage"):
        u = con.execute("SELECT * FROM eq_usage_daily WHERE page_id = ?", (src[r[L.SOURCE]],)).fetchone()
        sf = {k: state_of(con, u[f"{k}_field_id"]) for k in ("start", "end", "total")}
        for k, col in (("start", "계기 시작"), ("end", "계기 종료"), ("total", "계기 총")):
            assert r[col] == (u[f"meter_{k}"] if sf[k] == "value" else None), col
        for k, col in (("start", "시각 시작"), ("end", "시각 종료")):
            assert r[col] == (u[f"clock_{k}"] if sf[k] == "value" else None), col
        if site is not None:
            tpl = site.templates[u["source_form"]]
            roles = {reg["name"]: reg.get("role") for reg in tpl.regions}
            shifts = [f["field_id"] for f in con.execute("SELECT * FROM doc_field WHERE page_id = ?", (u["page_id"],))
                      if is_shift_cell(roles.get(f["region"]), f)]
            sh_ok = all_sure(state_of(con, f) for f in shifts)
            assert r["근무 분"] == (u["shift_minutes"] if sh_ok else None)
            basis = {"meter": ["start", "end"], "clock": ["start", "end"], "total": ["total", "start", "end"]}.get(u["hours_basis"])
            ok = all_sure(sf[k] for k in basis) if basis else (sh_ok if u["hours_basis"] == "shifts" else True)
            assert r["가동 시간"] == (u["hours"] if ok else None)
    # 계기 검산: 견준 칸 전부가 확정일 때만 두 값·차이
    for r in business_rows(b, "xcheck_usage"):
        kind = {v: k for k, v in L.CHECK_KIND.items()}[r["종류"]]
        rows = [x for x in con.execute("SELECT * FROM xcheck_usage WHERE page_id = ? AND check_kind = ?", (src[r[L.SOURCE]], kind))
                if x["field_a"] == r[L.FIELD_ID]]
        assert len(rows) == 1
        x = rows[0]
        ok = all_sure(state_of(con, f) for f in check_cells(con, site, x)) if site is not None else True
        assert (r["값 A"], r["값 B"], r["차이"]) == ((x["value_a"], x["value_b"], x["diff"]) if ok else (None, None, None))
        prov = not ok and x["result"] not in ("first", "unknown")
        assert r["결과"] == L.CHECK_RESULT[x["result"]] + (L.PROVISIONAL if prov else "")
    # 점검: 이상 유·무는 두 ✓ 칸(핸들러가 비운 칸은 확정된 빈 칸)이 확정일 때만, 점검내역은 비고 칸이 확정일 때만
    for r in business_rows(b, "inspection"):
        i = con.execute("SELECT i.*, p.template_name FROM insp_daily i JOIN doc_page p ON p.page_id = i.page_id "
                        "WHERE i.source_field_id = ?", (r[L.FIELD_ID],)).fetchone()
        rem = state_of(con, i["source_field_id"])
        assert r["점검내역"] == (i["remark"] or None if rem == "value" else None)
        if site is not None:
            tpl = site.templates[i["template_name"]]
            region, yes, no, _ = layout(tpl)
            row = int(i["source_field_id"].rsplit(":", 1)[1])
            fields = {f["field_id"]: dict(f) for f in con.execute("SELECT * FROM doc_field WHERE page_id = ?", (i["page_id"],))}
            over = get_handler(tpl.handler).export_cells(tpl, fields)[0]
            ids = [f"{i['page_id']}:{region}:{c}:{row}" for c in (yes, no)]
            ok = all_sure("empty" if over.get(f) == "empty" else state_of(con, f) for f in ids)
            assert r["이상"] == (L.ABNORMAL[i["abnormal"]] if ok else None)


# ── 모델 = DB ──────────────────────────────────────────────────────────────
def test_daily_model_matches_the_db_oracle(oracle_run, synth):
    con, site = oracle_run.con, oracle_run.site
    seen = Counter()
    for day in days_of(con):
        b = book(con, site, day)
        seen += check_form_cells(site, con, b)
        check_business(con, b, day, site)
        names = [s["name"] for s in b["sheets"]]
        assert names[0] == L.SUMMARY and all(len(n) <= 31 for n in names) and len(set(names)) == len(names)
    assert seen["value"] and seen["mark"] and seen["printed"] and seen["meta"] and seen["empty"]
    # 운반 값: 일보 시트의 확정된 횟수 = 정답 (oracle 은 숫자 칸을 자동 적재한다)
    truth = {(d["date"], h["source"], h["material"], h["level"], h["shift"]): h["trips"]
             for d in synth.truth["days"] for h in d["haul_log"]}
    got = {}
    for day in days_of(con):
        for r in business_rows(book(con, site, day), "haul_log"):
            if r["횟수"] is not None:
                got[(day, r[L.SOURCE], r["광종"], r["편"], {v: k for k, v in L.SHIFT.items()}[r["주야"]])] = r["횟수"]
    assert got == truth


def test_daily_model_of_usage_logs(usage_run):
    """가동 일보(인쇄 층, 동시 판): 칸의 상태와 업무 시트 — 하루 두 장인 장비는 두 행 그대로 (합치지 않는다)."""
    con, site = usage_run["pipe"].con, usage_run["pipe"].site
    for day in days_of(con):
        b = book(con, site, day)
        check_form_cells(site, con, b)
        check_business(con, b, day, site)
    # 가동 기록은 쪽 하나에 한 행 — 같은 날 같은 장비의 둘째 장(합성 첫날의 TRUCK)도 따로 한 행이다 (합치지 않는다)
    usage_pages = {t.name for t in site.templates.values() if t.handler == "usage"}
    for day in days_of(con):
        n = sum(1 for r in con.execute("SELECT template_name FROM doc_page WHERE work_date = ? AND status = 'loaded'", (day,))
                if r[0] in usage_pages)
        assert len(business_rows(book(con, site, day), "usage")) == n
    # 동시 판은 계열로 한 장 — 판 A·B 의 쪽이 같은 시트에 있는 날이 있다
    tn = dict(con.execute("SELECT page_id, template_name FROM doc_page"))
    both = False
    for day in days_of(con):
        for s in book(con, site, day)["sheets"]:
            names = {tn[pid] for pid in blocks(s)}
            assert len({site.templates[n].sig_family for n in names}) <= 1
            both |= {"synth_usage_log", "synth_usage_log_b"} <= names
    assert both


def test_summary_counts_match_direct_counts(oracle_run, usage_run):
    for pipe in (oracle_run, usage_run["pipe"]):
        con, site = pipe.con, pipe.site
        for day in days_of(con):
            b = book(con, site, day)
            rows = sheet(b, L.SUMMARY)["rows"]
            got = Counter()
            for r in rows:
                if len(r) == 3 and isinstance(r[2][0], int):
                    got[(r[0][0], r[1][0])] = r[2][0]
            for st, n in con.execute("SELECT status, COUNT(*) FROM doc_page WHERE work_date = ? GROUP BY 1", (day,)):
                assert got[(L.SEC_PAGES, L.PAGE_STATUS[st])] == n
            assert got[(L.SEC_PAGES, L.TOTAL)] == con.execute("SELECT COUNT(*) FROM doc_page WHERE work_date = ?", (day,)).fetchone()[0]
            for st, n in con.execute("SELECT status, COUNT(*) FROM xcheck_haul WHERE work_date = ? GROUP BY 1", (day,)):
                assert got[(L.SEC_XCHECK, L.XCHECK_STATUS[st])] == n
            for k, res, n in con.execute("SELECT check_kind, result, COUNT(*) FROM xcheck_usage WHERE work_date = ? GROUP BY 1, 2",
                                         (day,)):
                assert got[(L.SEC_XUSAGE, f"{L.CHECK_KIND[k]} — {L.CHECK_RESULT[res]}")] == n
            # 검수 대기 칸의 수: 요약 = 양식 시트의 ?·판독 불가 칸 = 쪽 머리의 수의 합
            forms = [s for s in b["sheets"][1:] if blocks(s)]
            n_cells = sum(1 for s in forms for row in s["rows"] for c in row if c[1] in UNSURE)
            heads = sum(row[5][0] for s in b["sheets"][1:] for row in s["rows"] if row and row[0] == [L.PAGE_HEAD, "h"])
            assert got[(L.SEC_PENDING, L.TOTAL)] == n_cells == heads


# ── 확정되지 않은 값은 싣지 않는다 ──────────────────────────────────────────────
class LowConfidence(OracleRecognizer):
    """정답을 읽되 신뢰도를 기준 아래로 낸다 — 값이 있는 검수 대기 칸 (기계 값이 value_final 에 들어 있다)."""
    name = "low"

    def recognize(self, crops, contexts):
        return [Recognition(r.text, 0.5 if r.text else 0.0, [], self.name) for r in super().recognize(crops, contexts)]


@pytest.fixture(scope="module")
def low_day(tmp_path_factory):
    """하루치(기본 묶음의 첫날과 같은 바이트)를 가짜 인식기로. 검수 파일은 root 아래."""
    root = tmp_path_factory.mktemp("low_day")
    syn = generate(root / "data", days=1, seed=0)
    st = Settings(site=syn.site, archive_root=syn.scans, work_root=root / "work", reviews=root / "기록" / "reviews.jsonl",
                  save_aligned=False)
    pipe = Pipeline(st, recognizer=LowConfidence(load_answers_json(syn.answers_path)))
    pipe.run([syn.scans])
    return {"pipe": pipe, "settings": st, "synth": syn, "root": root}


def machine_values_of(con) -> set:
    """검수 대기 칸의 기계 값 (value_final — 싣지 않아야 한다)."""
    return {r[0] for r in con.execute("SELECT value_final FROM doc_field WHERE review_status = 'pending' AND value_final <> ''")}


def test_unconfirmed_values_are_not_exported(null_run, low_day):
    # null: 값이 있는 수 칸은 전부 검수 대기
    con, site = null_run.con, null_run.site
    for day in days_of(con):
        b = book(con, site, day)
        for pid, region, row, name, (_value, style) in form_cells(site, con, b):
            f = con.execute("SELECT kind, format, has_value FROM doc_field WHERE field_id = ?", (f"{pid}:{region}:{name}:{row}",)).fetchone()
            if f["kind"] == "handwritten_number" and f["has_value"] and region != "fields":
                assert style == "pending", (pid, region, name)
    # 가짜 인식기: 검수 대기 칸에 기계 값이 들어 있다 — 양식 시트·업무 시트의 어느 칸에도 없다
    con, site = low_day["pipe"].con, low_day["pipe"].site
    pending = {r["field_id"]: r["value_final"] for r in con.execute(
        "SELECT field_id, value_final FROM doc_field WHERE review_status = 'pending' AND value_final <> ''")}
    assert len(pending) > 20
    day = days_of(con)[0]
    b = book(con, site, day)
    check_form_cells(site, con, b)
    check_business(con, b, day, site)                          # 점검내역(비고)도 확정일 때만 — 기계 값이 새지 않는다
    for pid, region, row, name, (value, style) in form_cells(site, con, b):
        if f"{pid}:{region}:{name}:{row}" in pending:
            assert style in UNSURE and value in (L.PENDING_MARK, L.ILLEGIBLE_MARK)
    machine = {typed(v, "integer") for v in pending.values()}
    for key in ("haul_log", "haul_matrix", "tally"):
        for r in business_rows(b, key):
            if r[L.STATE] in (L.ROW_STATE["pending"], L.ROW_STATE["illegible"]):
                assert r.get("횟수", r.get("수")) is None
    assert L.MACHINE not in [c[0] for s in b["sheets"] for c in (s["rows"][0] if s["rows"] else [])]
    for r in business_rows(b, "xcheck_haul"):                  # 교차검증의 횟수는 운반 행이 전부 확정일 때만
        if r[L.STATE] != L.ROW_STATE["value"]:
            assert r["일보 횟수"] is None and r["행렬 횟수"] is None and r["판정"].endswith(L.PROVISIONAL)
    # machine_values = true: 따로 둔 열에만, 확정이 아닌 행에만
    bm = book(con, site, day, machine_values=True)
    rows = business_rows(bm, "haul_log")
    with_machine = [r for r in rows if r[L.MACHINE] is not None]
    assert with_machine and all(r[L.STATE] == L.ROW_STATE["pending"] and r["횟수"] is None for r in with_machine)
    assert {r[L.MACHINE] for r in with_machine} <= machine


def test_monthly_long_table_has_no_unconfirmed_values(low_day):
    """월별의 긴 표(운반): 확정이 아닌 행의 횟수는 비고, 기계 값은 machine_values 일 때만 따로 둔 열에 (일별과 같은 규칙)."""
    from minedocscan.export.monthly import monthly_book

    con, site = low_day["pipe"].con, low_day["pipe"].site
    days = days_of(con)
    machine = {typed(v, "integer") for v in machine_values_of(con)}
    n = con.execute("SELECT COUNT(*) FROM prod_haul WHERE has_value = 1 OR review_status = 'pending'").fetchone()[0]
    for mv in (False, True):
        with read_txn(con):
            b = monthly_book(con, site, days[0][:7], days, machine_values=mv)
        s = sheet(b, L.LONG["haul"])
        names = [c[0] for c in s["rows"][0]]
        rows = [{k: c[0] for k, c in zip(names, r, strict=True)} for r in s["rows"][1:]]
        assert len(rows) == n
        unsure = [r for r in rows if r[L.STATE] in (L.ROW_STATE["pending"], L.ROW_STATE["illegible"])]
        assert len(unsure) > 20 and all(r["횟수"] is None for r in unsure)
        assert all(r["횟수"] is not None for r in rows if r[L.STATE] == L.ROW_STATE["value"])
        if not mv:
            assert L.MACHINE not in names
        else:
            with_machine = [r for r in rows if r[L.MACHINE] is not None]
            assert with_machine and all(r in unsure for r in with_machine)
            assert {r[L.MACHINE] for r in with_machine} <= machine


def test_reviews_turn_cells_into_values_and_illegible(low_day, tmp_path):
    from dataclasses import replace

    con = clone_db(low_day["pipe"].con)
    site, st = low_day["pipe"].site, replace(low_day["settings"], reviews=tmp_path / "reviews.jsonl")
    day = days_of(con)[0]
    a, c = (r["source_field_id"] for r in con.execute(
        "SELECT source_field_id FROM prod_haul WHERE work_date = ? AND source_role = 'log' AND review_status = 'pending' "
        "ORDER BY haul_id LIMIT 2", (day,)))
    save(con, site, st, review_from_field(con, a, "value", "7", "jp"))
    save(con, site, st, review_from_field(con, c, "illegible", "", "jp"))
    b = book(con, site, day)
    by_id = {r[L.FIELD_ID]: r for r in business_rows(b, "haul_log")}
    assert by_id[a]["횟수"] == 7 and by_id[a][L.STATE] == L.ROW_STATE["value"]
    assert by_id[c]["횟수"] is None and by_id[c][L.STATE] == L.ROW_STATE["illegible"]
    cells = {f"{pid}:{reg}:{n}:{row}": v for pid, reg, row, n, v in form_cells(site, con, b)}
    assert cells[a] == [7, "value"] and cells[c] == [L.ILLEGIBLE_MARK, "illegible"]


# ── 핸들러가 고쳐 말하는 칸 ─────────────────────────────────────────────────────
def test_inspection_unused_page_and_margin_rows(null_run):
    con, site = null_run.con, null_run.site
    tpl = site.templates["synth_inspection"]
    region, yes, no, _ = layout(tpl)
    margin = {r["row"] for r in tpl.region(region)["rows"] if not is_equipment_row(r)}
    assert margin
    unused_seen = used_seen = False
    for day in days_of(con):
        b = book(con, site, day)
        tn = dict(con.execute("SELECT page_id, template_name FROM doc_page"))
        s = next(s for s in b["sheets"] if any(tn[pid] == tpl.name for pid in blocks(s)))
        for pid, rows in blocks(s).items():
            marks = con.execute("SELECT has_value_raw FROM doc_field WHERE page_id = ? AND region = ? AND field_name IN (?, ?)",
                                (pid, region, yes, no)).fetchall()
            unused = all(m[0] is None for m in marks)
            notes = [r[1][0] for r in rows if len(r) == 2 and r[1][1] == "note"]
            assert notes == ([L_NO_MARKS()] if unused else [])
            for p, reg, row, name, (_value, style) in form_cells(site, con, {"sheets": [{"name": "x", "rows": rows}]}):
                if reg == region and name in (yes, no):
                    if unused or row in margin:
                        assert style != "pending", (p, row, name)       # 비어 있다 — ? 가 아니다
            unused_seen |= unused
            used_seen |= not unused
    assert unused_seen and used_seen                              # 합성 둘째 날은 점검하지 않은 쪽이다


def L_NO_MARKS():
    from minedocscan.handlers.inspection import NO_MARKS
    return NO_MARKS


@pytest.fixture(scope="module")
def usage_null_day(tmp_path_factory):
    root = tmp_path_factory.mktemp("usage_null_day")
    syn = generate(root / "data", days=1, seed=0, usage_only=True)
    pipe = Pipeline(Settings(site=syn.site, archive_root=syn.scans, work_root=root / "work", reviews=root / "reviews.jsonl",
                             save_aligned=False))
    pipe.run([syn.scans])
    return pipe


def is_activities(tpl, region: str) -> bool:
    return tpl.handler == "usage" and any(r["name"] == region and r.get("role") == "activities" for r in tpl.regions)


def pending_parts(site, con, day: str) -> tuple[int, Counter]:
    """적재된 쪽의 pending 필드를 나눈다: 엑셀에서 ?·판독 불가로 보이는 것(cells)과 아닌 것 (4.2 — 값이 있는 메타 필드, 여백 행의
    ✓ 칸, 점검하지 않은 쪽의 ✓ 칸, 작업 표의 글자 칸). 돌려주는 값: (pending 필드 수, 나눈 수)."""
    pend = [dict(r) for r in con.execute(
        "SELECT f.*, p.template_name FROM doc_field f JOIN doc_page p ON f.page_id = p.page_id "
        "WHERE p.work_date = ? AND p.status = 'loaded' AND f.review_status = 'pending'", (day,))]
    meta = {(r[0], r[1]) for r in con.execute(
        "SELECT m.page_id, m.meta_key FROM doc_page_meta m JOIN doc_page p ON m.page_id = p.page_id "
        "WHERE p.work_date = ? AND m.value IS NOT NULL", (day,))}
    parts = Counter()
    for f in pend:
        tpl = site.templates[f["template_name"]]
        mk = next((x.get("meta_key") for x in tpl.fields if x["name"] == f["field_name"]), None) if f["region"] == "fields" else None
        if mk and (f["page_id"], mk) in meta:
            parts["meta"] += 1
        elif tpl.handler == "inspection" and f["kind"] == "checkmark" and f["reviewed_by"] is None:
            region, yes, no, _ = layout(tpl)
            marks = con.execute("SELECT has_value_raw FROM doc_field WHERE page_id = ? AND region = ? AND field_name IN (?, ?)",
                                (f["page_id"], region, yes, no)).fetchall()
            eq_rows = {r["row"] for r in tpl.region(region)["rows"] if is_equipment_row(r)}
            if all(m[0] is None for m in marks):
                parts["unused"] += 1
            elif f["row_no"] not in eq_rows and f["has_value_raw"] != 1:
                parts["margin"] += 1
            else:
                parts["cells"] += 1
        elif f["kind"] == "handwritten_text" and f["reviewed_by"] is None and is_activities(tpl, f["region"]):
            parts["activities"] += 1
        else:
            parts["cells"] += 1
    return len(pend), parts


def test_activities_text_cells_and_the_pending_count(usage_null_day, usage_run, null_run):
    """작업 표의 검수 대기인 글자 칸은 ● (oracle 이면 값). 검수 대기 칸의 수와 DB 의 pending 필드 수의 관계 (4.2):
    pending 필드 = 검수 대기 칸 + 값이 있는 메타 필드 + 여백 행 ✓ + 점검하지 않은 쪽 ✓ + 작업 표의 글자 칸."""
    seen = Counter()
    for pipe, oracle in ((usage_null_day, False), (usage_run["pipe"], True), (null_run, False)):
        con, site = pipe.con, pipe.site
        tn = dict(con.execute("SELECT page_id, template_name FROM doc_page"))
        for day in days_of(con):
            b = book(con, site, day)
            styles = check_form_cells(site, con, b)
            for pid, reg, _row, _n, v in form_cells(site, con, b):
                tpl = site.templates[tn[pid]]
                if is_activities(tpl, reg) and v[1] != "printed":
                    seen[("oracle" if oracle else "null", v[1])] += 1
            n, parts = pending_parts(site, con, day)
            n_unsure = sum(styles[s] for s in UNSURE)
            assert parts["cells"] == n_unsure
            assert n == n_unsure + parts["meta"] + parts["margin"] + parts["unused"] + parts["activities"]
            seen.update({k: v for k, v in parts.items() if k != "cells"})
    assert seen[("null", "present")] and not seen[("null", "pending")] and seen[("oracle", "value")]
    assert seen["meta"] and seen["margin"] and seen["unused"] and seen["activities"]


# ── 파일 ───────────────────────────────────────────────────────────────────
def model_values(s: dict) -> list[list]:
    """모델의 시트를 되읽은 값과 같은 모양으로 (만든 시각·프로그램 판 칸은 빼고 견준다)."""
    rows = []
    for row in s["rows"]:
        vals = [None if c[1] in ("stamp_time", "stamp_version") else c[0] for c in row]
        while vals and vals[-1] is None:
            vals.pop()
        rows.append(vals)
    while rows and not rows[-1]:
        rows.pop()
    return rows


def test_the_file_equals_the_model(oracle_run, tmp_path):
    con, site = oracle_run.con, oracle_run.site
    out = tmp_path / "엑셀 폴더"
    out.mkdir()
    r = export_excel(con, site, out, full=True, made_at="2030-02-01T00:00:00Z")
    assert len(daily(r.written)) == len(days_of(con)) and not r.failed
    for day in days_of(con):
        b = book(con, site, day)
        got = read_values(out / daily_path(day))
        assert list(got) == [s["name"] for s in b["sheets"]]
        for s in b["sheets"]:
            want = model_values(s)
            have = got[s["name"]]
            stamp = [i for i, row in enumerate(s["rows"]) if any(c[1].startswith("stamp") for c in row)]
            assert all(have[i][1] for i in stamp)                # 만든 시각·판이 적혀 있다
            for i in stamp:
                have[i] = have[i][:1]
                want[i] = want[i][:1]
            assert have == want, s["name"]
        logs = got[L.SHEETS["haul_log"]]
        assert any(isinstance(v, int) for row in logs[1:] for v in row)            # 수는 수
        assert all(isinstance(row[0], str) for row in logs[1:])                     # 날짜는 글자
    # 표시: 바탕색과 글꼴, 틀 고정, 자동 필터 (모델의 표시대로)
    from openpyxl import load_workbook

    from minedocscan.export.xlsx import FILL

    day = days_of(con)[0]
    b = book(con, site, day)
    wb = load_workbook(str(out / daily_path(day)))
    seen = Counter()
    for s in b["sheets"]:
        ws = wb[s["name"]]
        assert ws.freeze_panes == (f"A{s['freeze'] + 1}" if s["freeze"] else None)
        assert bool(ws.auto_filter.ref) == bool(s["filter"])
        for i, row in enumerate(s["rows"], start=1):
            for j, (value, style) in enumerate(row, start=1):
                c = ws.cell(row=i, column=j)
                fills = [t for t in style.split() if t in FILL]
                if fills:
                    assert c.fill.start_color.rgb.endswith(FILL[fills[-1]]), (s["name"], i, j, style)
                    seen[fills[-1]] += 1
                if style == "h":
                    assert c.font.bold
                if isinstance(value, int | float) and not isinstance(value, bool):
                    assert c.data_type == "n" and c.value == value
                elif isinstance(value, str) and not style.startswith("stamp"):
                    assert c.data_type == "s" and c.value == value
    assert seen["printed"] and seen["meta"] and seen["h"]
    wb.close()


# ── 바뀐 것만 쓴다, 지운다, 쓰지 못해도 죽지 않는다 ──────────────────────────────
def test_only_changed_files_are_written_and_empty_days_are_removed(null_run, tmp_path):
    con = clone_db(null_run.con)
    site = null_run.site
    st = Settings(site=site.root, reviews=tmp_path / "reviews.jsonl")
    out = tmp_path / "out"
    out.mkdir()
    days = days_of(con)
    r1 = export_excel(con, site, out, full=True)
    assert daily(r1.written) == sorted(daily_path(d) for d in days)
    mtimes = {p: (out / p).stat().st_mtime_ns for p in r1.written}
    r2 = export_excel(con, site, out, full=True)
    assert r2.written == [] and r2.unchanged == len(r1.written)
    assert {p: (out / p).stat().st_mtime_ns for p in r1.written} == mtimes
    # 운반 횟수 칸 하나를 검수하면 그 날짜의 파일만
    fid, day = con.execute("SELECT source_field_id, work_date FROM prod_haul WHERE review_status = 'pending' "
                           "AND source_role = 'log' ORDER BY haul_id LIMIT 1").fetchone()
    save(con, site, st, review_from_field(con, fid, "value", "5", "jp"))
    inodes = {p: (out / p).stat().st_ino for p in r1.written}
    r3 = export_excel(con, site, out, full=True, made_at="2030-03-01T00:00:00Z")
    assert daily(r3.written) == [daily_path(day)]
    after = {p: (out / p).stat().st_ino for p in r1.written}
    assert {p for p in r1.written if after[p] != inodes[p]} == set(r3.written)       # 바꿔 넣었다 (제자리에 덮지 않는다)
    stamp = read_values(out / daily_path(day))[L.SUMMARY]
    assert [L.SEC_STAMP_TIME, "2030-03-01T00:00:00Z"] in stamp
    # 기록 파일을 지우면 전부 다시 쓴다 (지우지는 않는다)
    (out / RECORD_NAME).unlink()
    r4 = export_excel(con, site, out, days={days[0]})                  # 범위를 주어도 기록을 잃었으면 전부 훑는다
    assert daily(r4.written) == sorted(daily_path(d) for d in days) and not r4.deleted and r4.record_lost
    # 쪽이 없어진 날짜의 파일은 지운다 — 그 폴더에 따로 둔 파일은 남는다
    mine = out / daily_path(days[0]).rsplit("/", 1)[0] / "내 메모.xlsx"
    mine.write_bytes(b"x")
    con.execute("DELETE FROM doc_page WHERE work_date = ?", (days[0],))
    con.commit()
    r5 = export_excel(con, site, out, full=True)
    assert daily(r5.deleted) == [daily_path(days[0])] and not (out / daily_path(days[0])).exists() and mine.exists()
    # 날짜를 주면 그 날짜만 본다
    r6 = export_excel(con, site, out, days={days[1]})
    assert (r6.written, r6.unchanged, r6.deleted) == ([], 1, [])
    # 기록을 잃은 뒤의 전부 훑기: 쪽이 없는 날짜의 파일이 폴더에 있으면 지우지 않고 센다
    stale = out / daily_path("2029-12-31")
    stale.parent.mkdir(parents=True, exist_ok=True)
    stale.write_bytes(b"x")
    (out / RECORD_NAME).unlink()
    r7 = export_excel(con, site, out, full=True)
    assert r7.kept == 1 and stale.exists() and not r7.deleted


def test_a_file_that_cannot_be_replaced_is_counted_and_written_next_time(null_run, tmp_path, monkeypatch):
    import minedocscan.export.writer as w

    con, site = null_run.con, null_run.site
    out = tmp_path / "out"
    out.mkdir()
    day = days_of(con)[0]
    real = w._replace

    def locked(src, dst):
        raise PermissionError("열려 있다")

    monkeypatch.setattr(w, "_replace", locked)
    r = export_excel(con, site, out, days={day})                       # 기록이 없다 — 전부 훑는다
    assert daily_path(day) in r.failed and not r.written
    assert [p.name for p in out.rglob("*") if p.is_file()] == [RECORD_NAME]      # 임시 파일이 남지 않는다 (xlsx 도 없다)
    monkeypatch.setattr(w, "_replace", real)
    r = export_excel(con, site, out, days={day})
    assert daily_path(day) in r.written and not r.failed
    # 기록 파일을 읽지 못하면(잠겨 있다) 아무것도 쓰지 않는다 — 기록을 덮어쓰지 않는다
    before = (out / RECORD_NAME).read_bytes()
    monkeypatch.setattr(w, "load_record", lambda out: (_ for _ in ()).throw(w.RecordUnreadable("잠김")))
    r = export_excel(con, site, out, full=True)
    assert r.failed == [RECORD_NAME] and not r.written and (out / RECORD_NAME).read_bytes() == before
    monkeypatch.undo()
    # 엑셀 폴더가 없으면 만들지 않는다
    gone = tmp_path / "끊긴 폴더"
    r = export_excel(con, site, gone, full=True)
    assert r.missing_dir and not gone.exists()


def test_refusals_and_non_iso_dates(null_run, tmp_path):
    repo = Path(__file__).resolve().parents[1]
    with pytest.raises(ExportError):
        check_out_dir(repo / "out" / "xl")
    assert check_out_dir(repo / "out" / "xl", allow_in_repo=True)
    inbox, arch = tmp_path / "스캐너", tmp_path / "보관"
    with pytest.raises(ExportError):
        check_out_dir(inbox / "엑셀", Settings(inbox=inbox))
    with pytest.raises(ExportError):
        check_out_dir(inbox, Settings(inbox=inbox))
    with pytest.raises(ExportError):
        check_out_dir(arch / "intake" / "엑셀", Settings(archive_root=arch))
    ok = tmp_path / "엑셀"                                               # 옆 폴더는 된다 (이름이 앞부분만 같아도)
    assert check_out_dir(ok, Settings(inbox=inbox, archive_root=arch)) == ok
    assert check_out_dir(tmp_path / "스캐너2", Settings(inbox=inbox, archive_root=arch))
    con = clone_db(null_run.con)
    day = days_of(con)[0]
    con.execute("UPDATE doc_page SET work_date = '07/01/2030' WHERE work_date = ?", (day,))
    con.commit()
    out = tmp_path / "out"
    out.mkdir()
    r = export_excel(con, null_run.site, out, full=True)
    assert r.skipped_dates == 1 and not any("07/01" in p for p in r.written)
    files = {p.relative_to(out).as_posix() for p in out.rglob("*") if p.is_file()}
    assert files == {*r.written, RECORD_NAME}                              # 이름 규칙에 맞는 파일만 — 날짜가 경로가 되지 않았다
    assert daily(r.written) == sorted(daily_path(d) for d in days_of(con) if d != '07/01/2030')
    assert {p.relative_to(out).as_posix() for p in out.rglob("*") if p.is_dir()} <= {"daily", "daily/2030-01", "monthly"}
    # 바퀴 끝의 내보내기도 알린다: 바퀴를 시작할 때(어느 조각에도 들지 않는 날짜 — 한 번), 그 날짜를 건드린 바퀴, 그 달의 조각
    from minedocscan.export.auto import AutoExport
    from minedocscan.touched import Touched

    x = AutoExport(Settings(excel_dir=out, export_sweep_minutes=0.0), null_run.site, clock=lambda: 0.0)
    assert x.after_round(con, Touched()).skipped_dates == 1
    while x.sweep.running:
        assert x.after_round(con, Touched()).skipped_dates == 0
    assert x.after_round(con, Touched(dates={"07/01/2030"})).skipped_dates == 1
    con.execute("UPDATE doc_page SET work_date = '2030-01-3x' WHERE work_date = '07/01/2030'")
    con.commit()
    assert export_excel(con, null_run.site, out, slices=["2030-01"]).skipped_dates == 1


def test_export_command_refuses_while_the_pipeline_runs(null_run, tmp_path, capsys):
    from minedocscan.cli import main
    from minedocscan.pipeline.lock import PipelineLock

    s = null_run.settings
    out = tmp_path / "out"
    out.mkdir()
    lock = PipelineLock.for_settings(s).acquire()
    try:
        with pytest.raises(SystemExit) as e:
            main(["export", "excel", str(out), "--site", str(s.site), "--work-root", str(s.work_root)])
        assert "이미 돌고" in str(e.value)
    finally:
        lock.release()
    day = days_of(null_run.con)[0]
    assert main(["export", "excel", str(out), "--site", str(s.site), "--work-root", str(s.work_root), "--date", day]) == 0
    assert daily_path(day) in capsys.readouterr().out
    assert main(["export", "excel", str(out), "--site", str(s.site), "--work-root", str(s.work_root), "--date", day]) == 0
    assert "쓴 파일 0, 그대로 " in capsys.readouterr().out


# ── 점검표의 모름, generic 양식의 ✓ ───────────────────────────────────────────
def test_unknown_inspection_rows_have_no_abnormal(null_run, tmp_path):
    con = clone_db(null_run.con)
    site = null_run.site
    st = Settings(site=site.root, reviews=tmp_path / "reviews.jsonl")
    tpl = site.templates["synth_inspection"]
    region, yes, no, _ = layout(tpl)
    i = con.execute("SELECT * FROM insp_daily WHERE abnormal IS NOT NULL ORDER BY inspection_id LIMIT 1").fetchone()
    row = int(i["source_field_id"].rsplit(":", 1)[1])                  # 비고 칸의 행 = 그 장비 행
    for c in (yes, no):
        save(con, site, st, review_from_field(con, f"{i['page_id']}:{region}:{c}:{row}", "illegible", "", "jp"))
    assert con.execute("SELECT abnormal FROM insp_daily WHERE inspection_id = ?", (i["inspection_id"],)).fetchone()[0] is not None
    rows = business_rows(book(con, site, i["inspection_date"]), "inspection")
    r = next(r for r in rows if r[L.FIELD_ID] == i["source_field_id"])
    assert r["이상"] is None and r[L.STATE] == L.ROW_STATE["illegible"]


def test_checkmark_cell_of_a_generic_template_shows_the_mark(tmp_path):
    """점검표가 아닌 양식의 ✓ 칸(값 없이 유무만)은 ✓, 서명은 ●. field_cell 이 4.2 의 첫째 표대로."""
    f = {"kind": "checkmark", "review_status": "auto", "reviewed_by": None, "has_value": 1, "value_final": None, "format": None}
    assert field_cell(f) == [L.CHECK_MARK, "mark"]
    assert field_cell({**f, "kind": "signature"}) == [L.PRESENT_MARK, "present"]
    assert field_cell({**f, "kind": "handwritten_number", "value_final": "12", "format": "integer"}) == [12, "value"]
    assert field_cell({**f, "kind": "handwritten_number", "value_final": "1234.5", "format": "reading"}) == [1234.5, "value"]
    assert field_cell({**f, "kind": "handwritten_text", "value_final": "08:00", "format": "time"}) == ["08:00", "value"]
    assert field_cell({**f, "has_value": 0}) == [None, "empty"]
    assert field_cell({**f, "review_status": "pending", "value_final": "9"}) == [L.PENDING_MARK, "pending"]
    assert field_cell({**f, "review_status": "pending", "reviewed_by": "jp", "value_final": "9"}) == [L.ILLEGIBLE_MARK, "illegible"]
    assert field_cell({**f, "kind": "printed", "value_final": "ORE"}, "x") == ["ORE", "printed"]
    assert field_cell({**f, "kind": "handwritten_text", "review_status": "pending"}, "V-101") == ["V-101", "meta"]
    # generic 핸들러의 작은 양식 하나로 쪽 블록까지
    from minedocscan.export.model import Pages, page_block

    tdir = tmp_path / "templates" / "t_generic"
    tdir.mkdir(parents=True)
    spec = {"name": "t_generic", "title": "작은 양식", "reference_image": "reference.png", "handler": "generic",
            "regions": [{"name": "main", "grid": {"ys": [0, 50, 100], "xs": [0, 100, 200]}, "header_rows": 1,
                         "columns": [{"idx": 0, "name": "what", "kind": "printed"}, {"idx": 1, "name": "done", "kind": "checkmark"}],
                         "rows": [{"row": 0, "key": "a", "what": "물 주기"}]}], "fields": []}
    (tdir / "template.yaml").write_text(yaml.safe_dump(spec, allow_unicode=True), encoding="utf-8")
    tpl = Template(tdir / "template.yaml")

    class Site:
        templates = {"t_generic": tpl}

    pid = "d-p1"
    fields = {f"{pid}:main:what:0": {**f, "field_id": f"{pid}:main:what:0", "kind": "printed", "value_final": "물 주기",
                                      "region": "main", "row_no": 0, "field_name": "what"},
              f"{pid}:main:done:0": {**f, "field_id": f"{pid}:main:done:0", "region": "main", "row_no": 0, "field_name": "done"}}
    page = {"page_id": pid, "template_name": "t_generic", "source_name": "s", "page_no": 1, "status": "loaded"}
    rows, n = page_block(Site, page, Pages([page], fields, {pid: fields}, {}))
    assert n == 0 and rows[3] == [["a", "label"], ["물 주기", "printed"], [L.CHECK_MARK, "mark"]]


# ── 표시 이름 ───────────────────────────────────────────────────────────────
@pytest.mark.slow                                  # 합성 묶음을 새로 만든다 — 기본 시험 시간 (tasks/0008 6절)
def test_display_names_show_in_the_sheets_and_change_nothing_else(null_run, tmp_path):
    syn = generate(tmp_path / "data", days=1, seed=0, display_names=True)
    pipe = Pipeline(Settings(site=syn.site, archive_root=syn.scans, work_root=tmp_path / "work", save_aligned=False))
    pipe.run([syn.scans])
    con, site = pipe.con, pipe.site
    day = days_of(con)[0]
    b = book(con, site, day)
    names = [s["name"] for s in b["sheets"]]
    assert "덤프트럭 운반 일보" in names and "일일 장비 점검표" in names
    check_form_cells(site, con, b)                            # 머리글·행 이름이 display (form_cells 가 display_of 로 본다)
    log = next(s for s in b["sheets"] if s["name"] == "덤프트럭 운반 일보")
    heads = [c[0] for r in log["rows"] for c in r if c[1] == "h"]
    assert "주간" in heads and "광종" in heads and any(r and r[0][0] == "광석 L0" for r in log["rows"])
    # doc_field(인쇄된 값 포함)·업무 테이블이 display 없이 돌린 것과 같다 (첫날은 바이트까지 같은 PDF — 같은 문서 ID)
    doc = con.execute("SELECT document_id FROM doc_document").fetchone()[0]
    skip = {"created_at", "received_at", "work_requested", "work_done", "source_path", "aligned_image"}

    def rows_of(c: sqlite3.Connection, table: str, where: str, arg) -> list:
        cols = [r[1] for r in c.execute(f"PRAGMA table_info({table})") if r[1] not in skip]
        return sorted(tuple(r) for r in c.execute(f"SELECT {', '.join(cols)} FROM {table} WHERE {where}", (arg,)))

    for t, where, arg in (("doc_field", "page_id LIKE ?", f"{doc}-%"), ("doc_page_meta", "page_id LIKE ?", f"{doc}-%"),
                          ("prod_haul", "work_date = ?", day), ("insp_daily", "inspection_date = ?", day),
                          ("xcheck_haul", "work_date = ?", day), ("eq_assignment_obs", "work_date = ?", day)):
        assert rows_of(con, t, where, arg) == rows_of(null_run.con, t, where, arg), t
    # 표시 이름이 없는 템플릿은 name·title
    plain = book(null_run.con, null_run.site, day)
    assert null_run.site.templates["synth_haul_log"].title in [s["name"] for s in plain["sheets"]] or \
        null_run.site.templates["synth_haul_log"].title[:31] in [s["name"] for s in plain["sheets"]]


def test_display_names_change_only_template_yaml(tmp_path):
    import hashlib

    a = generate(tmp_path / "a", days=1, seed=0)
    b = generate(tmp_path / "b", days=1, seed=0, display_names=True)

    def digest(root):
        return {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(root.rglob("*")) if p.is_file()}

    da, db = digest(a.root), digest(b.root)
    assert da.keys() == db.keys()
    assert {k for k in da if da[k] != db[k]} == {k for k in da if k.endswith("template.yaml")}


def test_template_check_rejects_a_display_column_and_bad_display(tmp_path, site):
    from minedocscan.tools.tpltools import check_template

    src = site.templates["synth_haul_log"].dir
    tdir = tmp_path / "t"
    tdir.mkdir()
    (tdir / "reference.png").write_bytes((src / "reference.png").read_bytes())
    spec = yaml.safe_load((src / "template.yaml").read_text(encoding="utf-8"))
    spec["regions"][0]["columns"][0]["name"] = "display"
    spec["fields"][0]["display"] = ""
    (tdir / "template.yaml").write_text(yaml.safe_dump(spec, allow_unicode=True), encoding="utf-8")
    errs = check_template(tdir)
    assert any("열 이름 'display'" in e for e in errs) and any("display 는 빈 문자열이 아닌" in e for e in errs)
    # 동시 판은 양식의 display 도 같아야 한다
    from minedocscan.forms.sitepack import variant_key_diff

    t = Template(src / "template.yaml")
    spec2 = yaml.safe_load((src / "template.yaml").read_text(encoding="utf-8"))
    spec2["display"] = "다른 이름"
    (tdir / "template.yaml").write_text(yaml.safe_dump(spec2, allow_unicode=True), encoding="utf-8")
    assert variant_key_diff(t, Template(tdir / "template.yaml")) == "양식의 display"


@pytest.mark.slow                                  # 합성 묶음을 새로 만든다 — 기본 시험 시간 (tasks/0008 6절)
def test_site_pack_loads_with_display(tmp_path):
    syn = generate(tmp_path / "d", days=1, seed=0, display_names=True, usage_logs=True, usage_variants=True)
    site = SitePack(syn.site)                                  # 동시 판의 display 가 같아 사이트 팩이 열린다
    assert site.templates["synth_usage_log_b"].display == site.templates["synth_usage_log"].display == "중기 운행일보"


# ── 검토에서 더한 것: 핸들러가 비우지 않는 칸, 요약의 수, 시트 이름, 수식, 순서 ───────────────────
def test_inspection_overrides_leave_unsure_marks_visible(null_run):
    """점검표의 ✓ 칸을 비우는 것은 '기계가 보지 못한 칸'뿐이다: 점검한 쪽의 장비 행 ?, 여백 행인데 기계가 표시를 본 ?,
    사람이 읽지 못한다고 한 칸(점검하지 않은 쪽이어도)은 그대로 보인다."""
    con = clone_db(null_run.con)
    site = null_run.site
    tpl = site.templates["synth_inspection"]
    region, yes, no, _ = layout(tpl)
    eq_rows = sorted(r["row"] for r in tpl.region(region)["rows"] if is_equipment_row(r))
    margin = sorted(r["row"] for r in tpl.region(region)["rows"] if not is_equipment_row(r))
    pages = [r[0] for r in con.execute("SELECT page_id FROM doc_page WHERE template_name = ? AND status = 'loaded' ORDER BY page_id",
                                       (tpl.name,))]

    def unused(pid):
        return all(m[0] is None for m in con.execute(
            "SELECT has_value_raw FROM doc_field WHERE page_id = ? AND region = ? AND field_name IN (?, ?)", (pid, region, yes, no)))

    used = next(p for p in pages if not unused(p))
    idle = next(p for p in pages if unused(p))
    a = f"{used}:{region}:{yes}:{eq_rows[0]}"                   # 점검한 쪽, 장비 행 — 기계가 정하지 못했다
    m = f"{used}:{region}:{yes}:{margin[0]}"                    # 점검한 쪽, 여백 행 — 기계가 표시를 보았다
    i = f"{idle}:{region}:{no}:{eq_rows[0]}"                     # 점검하지 않은 쪽 — 사람이 읽지 못한다고 했다
    con.execute("UPDATE doc_field SET review_status = 'pending', reviewed_by = NULL WHERE field_id = ?", (a,))
    con.execute("UPDATE doc_field SET review_status = 'pending', reviewed_by = NULL, has_value_raw = 1 WHERE field_id = ?", (m,))
    con.execute("UPDATE doc_field SET review_status = 'pending', reviewed_by = 'jp' WHERE field_id = ?", (i,))
    con.commit()
    cells = {}
    for day in {r[0] for r in con.execute("SELECT work_date FROM doc_page WHERE page_id IN (?, ?)", (used, idle))}:
        b = book(con, site, day)
        cells.update({f"{pid}:{reg}:{n}:{row}": v for pid, reg, row, n, v in form_cells(site, con, b)})
    assert cells[a] == [L.PENDING_MARK, "pending"]
    assert cells[m] == [L.PENDING_MARK, "pending"]
    assert cells[i] == [L.ILLEGIBLE_MARK, "illegible"]
    # 점검하지 않은 쪽이어도 사람이 표시를 확인한 칸이 있으면 "점검 표시 없음" 줄을 쓰지 않는다
    def note_of(pid):
        b = book(con, site, con.execute("SELECT work_date FROM doc_page WHERE page_id = ?", (pid,)).fetchone()[0])
        rows = next(blocks(s)[pid] for s in b["sheets"] if pid in blocks(s))
        return [r[1][0] for r in rows if len(r) == 2 and r[1][1] == "note"]

    assert note_of(idle) == [L_NO_MARKS()]
    con.execute("UPDATE doc_field SET review_status = 'reviewed', reviewed_by = 'jp', has_value = 1, value_final = '1' "
                "WHERE field_id = ?", (f"{idle}:{region}:{yes}:{eq_rows[0]}",))
    con.commit()
    assert note_of(idle) == []


def test_summary_page_counts_with_every_status_and_a_waiting_document(null_run):
    con = clone_db(null_run.con)
    site = null_run.site
    day = days_of(con)[0]
    pids = [r[0] for r in con.execute("SELECT page_id FROM doc_page WHERE work_date = ? ORDER BY page_id", (day,))]
    for pid, st in zip(pids, ("blank", "unknown_form", "duplicate", "discarded"), strict=False):
        con.execute("UPDATE doc_page SET status = ? WHERE page_id = ?", (st, pid))
    doc = con.execute("SELECT document_id FROM doc_page WHERE page_id = ?", (pids[0],)).fetchone()[0]
    con.execute("UPDATE doc_document SET work_requested = work_done + 1 WHERE document_id = ?", (doc,))
    con.commit()
    b = book(con, site, day)
    rows = sheet(b, L.SUMMARY)["rows"]
    pages = {r[1][0]: r[2][0] for r in rows if len(r) == 3 and r[0][0] == L.SEC_PAGES}
    want = {L.PAGE_STATUS[s]: n for s, n in con.execute("SELECT status, COUNT(*) FROM doc_page WHERE work_date = ? GROUP BY 1", (day,))}
    assert pages == {**want, L.TOTAL: len(pids)} and len(want) == 5
    waiting = next(r for r in rows if r and r[0][0] == L.SEC_WAITING)
    assert waiting[2][0] == 1 and waiting[1] == [L.WAITING_NOTE, "note"]
    forms = {r[1][0]: r[2][0] for r in rows if len(r) == 3 and r[0][0] == L.SEC_FORMS}
    assert sum(forms.values()) == pages[L.PAGE_STATUS["loaded"]]
    shown = {pid for s in b["sheets"] for pid in blocks(s)}               # 적재되지 않은 쪽은 양식 시트에 없다
    assert shown == {p for p in pids if con.execute("SELECT status FROM doc_page WHERE page_id = ?", (p,)).fetchone()[0] == "loaded"}


def test_sheet_names_are_valid_and_unique():
    from minedocscan.export.model import finish, sheet_names

    names = sheet_names(["요약", "a/b:c*d?[e]\\f", "가" * 40, "History", "HISTORY", "'따옴'", "", "운반", "운반", "운반".upper(),
                         "\x07종\x1f"])
    assert names[1] == "abcdef" and len(names[2]) == 31 and names[5] == "따옴" and names[6] == "시트" and names[10] == "종"
    assert names[3] == "History (2)" and names[4] == "HISTORY (3)"
    assert names[7] == "운반" and names[8] == "운반 (2)"
    assert len({n.casefold() for n in names}) == len(names) and all(0 < len(n) <= 31 for n in names)
    assert all(ch not in n for n in names for ch in "[]:*?/\\")
    long = sheet_names(["나" * 31, "나" * 31])
    assert long[1] == "나" * 27 + " (2)" and len(long[1]) == 31
    # 양식의 표시 이름이 업무 시트의 이름과 같아도 업무 시트의 이름은 그대로 — 양식 시트가 (2)
    b = finish("daily", "2030-01-07", [{"name": L.SUMMARY, "rows": []}, {"name": L.SHEETS["haul_log"], "rows": [], "family": "x"},
                                       {"name": L.SHEETS["haul_log"], "rows": []}])
    assert [s["name"] for s in b["sheets"]] == [L.SUMMARY, f"{L.SHEETS['haul_log']} (2)", L.SHEETS["haul_log"]]
    assert all("family" not in s for s in b["sheets"])


def test_text_that_looks_like_a_formula_stays_text(tmp_path):
    from openpyxl import load_workbook

    from minedocscan.export.model import cell, finish
    from minedocscan.export.xlsx import write_book

    tricky = ["=HYPERLINK(\"http://x\",\"y\")", "\x01=SUM(A1:A2)", "#N/A", "#DIV/0!", "+1", "-1", "@x", "=1+1"]
    b = finish("daily", "2030-01-07", [{"name": "t", "rows": [[cell(v, "value")] for v in tricky] + [[cell(3, "value")]]}])
    p = tmp_path / "t.xlsx"
    write_book(b, p, "2030-01-01T00:00:00Z", "v")
    ws = load_workbook(str(p))["t"]
    got = [ws.cell(row=i, column=1) for i in range(1, len(tricky) + 2)]
    assert all(c.data_type == "s" for c in got[:-1]) and got[-1].data_type == "n"
    assert [c.value for c in got[:-1]] == [v.replace("\x01", "") for v in tricky]
    assert read_values(p)["t"] == [[v.replace("\x01", "")] for v in tricky] + [[3]]


def test_meta_fields_fall_through_and_illegible_machine_values_are_hidden(null_run):
    """표 밖 필드: 쪽 메타 값이 있으면 그 값, 없으면 칸의 상태. 사람이 읽지 못한다고 한 칸에 기계가 자동 적재한 메타 값은 싣지 않는다
    (사람·파일명·결정의 값은 싣는다)."""
    from minedocscan.export.model import load_pages, page_block

    con, site = null_run.con, null_run.site
    day = days_of(con)[0]
    with read_txn(con):
        pages = load_pages(con, [day])
    p = next(p for p in pages.loaded if any(f.get("meta_key") for f in site.templates[p["template_name"]].fields))
    tpl = site.templates[p["template_name"]]
    f = next(f for f in tpl.fields if f.get("meta_key"))
    fid, key, pid = f"{p['page_id']}:fields:{f['name']}:-1", f["meta_key"], p["page_id"]
    row = pages.by_page[pid][fid]

    def shown(meta, source, **change):
        pages.meta[pid] = {key: meta} if meta is not None else {}
        pages.meta_source[pid] = {key: source} if meta is not None else {}
        pages.by_page[pid][fid] = {**row, **change}
        rows, _n = page_block(site, p, pages)
        return next(r[1] for r in rows if r and r[0] == [display_of(f, f["name"]), "label"])

    assert shown("V-101", "label") == ["V-101", "meta"]
    assert shown("V-101", "machine", review_status="pending", reviewed_by="jp") == [L.ILLEGIBLE_MARK, "illegible"]
    assert shown("V-101", "label", review_status="pending", reviewed_by="jp") == ["V-101", "meta"]
    assert shown("V-101", "machine", review_status="auto", reviewed_by=None) == ["V-101", "meta"]
    assert shown(None, None, review_status="pending", reviewed_by=None) == [L.PENDING_MARK, "pending"]
    assert shown(None, None, review_status="auto", reviewed_by=None, has_value=0) == [None, "empty"]


def reversed_db(con) -> sqlite3.Connection:
    """같은 행을 거꾸로 넣은 DB (행이 들어간 순서에 기대지 않는가 — CLAUDE.md '순서에 기대지 않는다')."""
    out = sqlite3.connect(":memory:")
    out.row_factory = sqlite3.Row
    tables = [r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type = 'table' AND name NOT LIKE 'sqlite_%'")]
    for (sql,) in con.execute("SELECT sql FROM sqlite_master WHERE sql IS NOT NULL AND name NOT LIKE 'sqlite_%' ORDER BY type DESC"):
        out.execute(sql)
    for t in tables:
        rows = con.execute(f"SELECT * FROM {t} ORDER BY rowid DESC").fetchall()
        if rows:
            out.executemany(f"INSERT INTO {t} VALUES ({','.join('?' * len(rows[0]))})", [tuple(r) for r in rows])
    out.commit()
    return out


def test_the_model_does_not_depend_on_row_order(oracle_run):
    """같은 DB 를 행의 순서만 바꿔 만들어도 책(해시)이 같다. 쪽 블록은 쪽의 순서(store/order.py)대로."""
    from minedocscan.export.model import book_hash
    from minedocscan.store.order import page_key

    con, site = oracle_run.con, oracle_run.site
    rev = reversed_db(con)
    order = "SELECT page_id FROM doc_page ORDER BY rowid"
    assert [r[0] for r in rev.execute(order)] == [r[0] for r in con.execute(order)][::-1]
    for day in days_of(con):
        a, b = book(con, site, day), book(rev, site, day)
        assert book_hash(a) == book_hash(b)
        pages = {r["page_id"]: dict(r) for r in con.execute(
            "SELECT p.*, d.source_rel, d.source_path FROM doc_page p JOIN doc_document d ON p.document_id = d.document_id "
            "WHERE p.work_date = ?", (day,))}
        for s in a["sheets"]:
            ids = list(blocks(s))
            assert ids == sorted(ids, key=lambda i: page_key(pages[i]["source_rel"], pages[i]["source_path"],
                                                              pages[i]["document_id"], pages[i]["page_no"]))


def test_confirmed_usage_values_after_review(usage_run, usage_synth, tmp_path):
    """계기·근무 시각 칸을 정답대로 검수하면 가동 기록의 값·가동 시간·검산의 두 값이 실린다 (확정이 아니면 비는 열이 채워진다)."""
    from conftest import review_usage

    con = clone_db(usage_run["pipe"].con)
    site = usage_run["pipe"].site
    st = Settings(site=site.root, reviews=tmp_path / "reviews.jsonl")
    review_usage(con, site, st, load_answers_json(usage_synth.answers_path), regions=("meter", "shifts"))
    filled = Counter()
    for day in days_of(con):
        b = book(con, site, day)
        check_business(con, b, day, site)
        for r in business_rows(b, "usage"):
            filled["hours"] += r["가동 시간"] is not None
            filled["meter"] += r["계기 시작"] is not None
            filled["clock"] += r["시각 시작"] is not None
        for r in business_rows(b, "xcheck_usage"):
            filled["xcheck"] += r["값 A"] is not None
    assert filled["hours"] and filled["meter"] and filled["clock"] and filled["xcheck"]


def test_one_broken_file_does_not_stop_the_others(null_run, tmp_path, monkeypatch):
    """한 날짜의 모델이 실패해도(템플릿이 DB 와 어긋났다 …) 다른 파일과 기록은 쓰고, 그 파일은 "쓰지 못함"으로 센다."""
    import minedocscan.export.writer as w

    con, site = null_run.con, null_run.site
    out = tmp_path / "out"
    out.mkdir()
    days = days_of(con)
    real = w.daily_book

    def broken(con_, site_, day, *a, **kw):
        if day == days[1]:
            raise KeyError("없는 표")
        return real(con_, site_, day, *a, **kw)

    monkeypatch.setattr(w, "daily_book", broken)
    r = export_excel(con, site, out, full=True)
    assert r.failed == [daily_path(days[1])] and daily_path(days[0]) in r.written and daily_path(days[2]) in r.written
    assert (out / RECORD_NAME).is_file()
    monkeypatch.setattr(w, "daily_book", real)
    r = export_excel(con, site, out, full=True)
    assert daily(r.written) == [daily_path(days[1])] and not r.failed


def test_template_drift_in_meter_checks_is_not_confirmed(usage_run, monkeypatch):
    """계기 검산이 견준 칸을 템플릿에서 찾지 못하면(표 이름이 바뀌었다) 값을 싣지 않는다 — 내보내기가 죽지도 않는다."""
    import minedocscan.export.business as biz

    con, site = usage_run["pipe"].con, usage_run["pipe"].site

    def gone(*a, **kw):
        raise KeyError("tally")

    monkeypatch.setattr(biz, "check_cells", gone)
    rows = [r for day in days_of(con) for r in business_rows(book(con, site, day), "xcheck_usage")]
    assert rows and len(rows) == con.execute("SELECT COUNT(*) FROM xcheck_usage").fetchone()[0]    # 볼 행이 있다 (tasks/0009 4.1 자)
    for r in rows:
        assert (r["값 A"], r["값 B"], r["차이"]) == (None, None, None)


def test_a_machine_meta_value_marked_illegible_is_hidden_in_business_sheets(null_run):
    con = clone_db(null_run.con)
    site = null_run.site
    h = con.execute("SELECT h.page_id, h.work_date, h.vehicle_no FROM prod_haul h JOIN doc_page_meta m ON m.page_id = h.page_id "
                    "AND m.meta_key = 'vehicle_no' WHERE h.source_role = 'log' AND h.vehicle_no IS NOT NULL LIMIT 1").fetchone()
    pid, day = h["page_id"], h["work_date"]
    shown = {r[L.SOURCE] for r in business_rows(book(con, site, day), "haul_log") if r["차량번호"] is not None}
    tpl = site.templates[con.execute("SELECT template_name FROM doc_page WHERE page_id = ?", (pid,)).fetchone()[0]]
    name = next(f["name"] for f in tpl.fields if f.get("meta_key") == "vehicle_no")
    con.execute("UPDATE doc_page_meta SET source = 'machine' WHERE page_id = ? AND meta_key = 'vehicle_no'", (pid,))
    con.execute("UPDATE doc_field SET review_status = 'pending', reviewed_by = 'jp' WHERE field_id = ?", (f"{pid}:fields:{name}:-1",))
    con.commit()
    src = next(r[L.SOURCE] for r in business_rows(book(con, site, day), "haul_log") if r[L.FIELD_ID].startswith(pid))
    rows = [r for r in business_rows(book(con, site, day), "haul_log") if r[L.SOURCE] == src]
    assert src in shown and rows and all(r["차량번호"] is None for r in rows)


# ── 기록 파일의 글자, 사본의 주인, 빈 작업 DB, 날짜 고르기 (tasks/0009 4.1 나·다·사·자) ─────────────────
OTHER_SITE = "OTHER-SITE-X"                        # 다른 사이트 팩의 이름 (합성 — 글에 찍히지 않아야 한다)


def snapshot(out: Path) -> dict[str, tuple[bytes, int]]:
    """폴더의 파일 전부(기록 파일 포함): 경로 → (바이트, 수정 시각)."""
    return {p.relative_to(out).as_posix(): (p.read_bytes(), p.stat().st_mtime_ns) for p in sorted(out.rglob("*")) if p.is_file()}


def record_of(out: Path) -> dict:
    """기록 파일을 UTF-8 JSON 으로 읽는다 (쓰는 것은 늘 UTF-8)."""
    import json

    return json.loads((out / RECORD_NAME).read_bytes().decode("utf-8"))


def all_files(con) -> list[str]:
    days = days_of(con)
    return sorted([daily_path(d) for d in days] + [monthly_path(m) for m in {d[:7] for d in days}])


def test_a_record_that_is_not_utf8_json_is_lost_and_a_locked_record_stops_the_export(null_run, tmp_path, monkeypatch):
    """기록 파일을 다른 인코딩으로 다시 저장했다(UTF-16)·쓰레기 바이트·빈 파일 → 깨졌다: 범위를 주어도 전부 다시 쓰고 기록을 UTF-8 로
    다시 적는다. 사이트 자리가 글자가 아니거나 빈 기록도 깨진 것이다. 있는데 읽지 못하는 기록(잠겨 있다) → 아무것도 쓰지 않고
    기록도 그대로 (4.1 나)."""
    import json

    con, site = null_run.con, null_run.site
    out = tmp_path / "out"
    out.mkdir()
    export_excel(con, site, out, full=True)
    rec = out / RECORD_NAME
    good = record_of(out)
    assert good["site"] == "synthetic" and sorted(good["files"]) == all_files(con)
    day = days_of(con)[0]
    odd_site = [json.dumps(dict(good, site=v)).encode() for v in (123, "", "  ", None)]   # 사이트 자리가 글자가 아니다·비었다
    for raw in (json.dumps(good, ensure_ascii=False).encode("utf-16"), b"\xff\xfe\x00{\x80 garbage \xc3", b"", *odd_site[:3]):
        rec.write_bytes(raw)
        r = export_excel(con, site, out, days={day}, months={day[:7]})
        assert r.record_lost and sorted(r.written) == all_files(con) and not r.deleted and not r.failed, raw[:8]
        again = record_of(out)                                           # UTF-8 JSON 으로 다시 적었다
        assert again["site"] == "synthetic" and sorted(again["files"]) == all_files(con)
        assert {k: v["sha"] for k, v in again["files"].items()} == {k: v["sha"] for k, v in good["files"].items()}
    rec.write_bytes(odd_site[3])                                         # "site": null 은 옛 기록(키가 없다)과 같다 — 받아들이고 적는다
    r = export_excel(con, site, out, days={day}, months={day[:7]})
    assert not r.record_lost and not r.other_site and not r.written and record_of(out)["site"] == "synthetic"
    # 있는데 읽지 못한다 (기록 파일만 — 엑셀 파일·DB 는 읽힌다): 이번에는 하지 않는다 — 기록을 덮어쓰지 않는다
    before = snapshot(out)
    real = Path.read_bytes

    def locked(self):
        if self.name == RECORD_NAME:
            raise PermissionError("잠겨 있다")
        return real(self)

    monkeypatch.setattr(Path, "read_bytes", locked)
    r = export_excel(con, site, out, full=True)
    monkeypatch.undo()
    assert r.failed == [RECORD_NAME] and not r.written and not r.deleted and not r.record_lost
    assert snapshot(out) == before


def test_a_folder_of_another_site_is_neither_written_nor_cleaned(null_run, tmp_path, capsys):
    """기록 파일의 사이트가 다른 이름이면 쓸 것(검수한 칸)도 지울 것(쪽이 없어진 날짜)도 그대로 둔다. 명령은 종료 코드 1 과 한 줄 —
    두 사이트의 이름 없이 (4.1 다)."""
    import json

    from minedocscan.cli import main

    con = clone_db(null_run.con)
    site = null_run.site
    st = Settings(site=site.root, reviews=tmp_path / "reviews.jsonl")
    out = tmp_path / "out"
    out.mkdir()
    export_excel(con, site, out, full=True)
    data = record_of(out)
    data["site"] = OTHER_SITE
    (out / RECORD_NAME).write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    days = days_of(con)
    fid = con.execute("SELECT source_field_id FROM prod_haul WHERE work_date = ? AND source_role = 'log' AND "
                      "review_status = 'pending' ORDER BY haul_id LIMIT 1", (days[1],)).fetchone()[0]
    save(con, site, st, review_from_field(con, fid, "value", "5", "jp"))             # 쓸 파일
    con.execute("DELETE FROM doc_page WHERE work_date = ?", (days[0],))               # 지울 파일
    con.commit()
    before = snapshot(out)
    for kw in ({"full": True}, {"days": {days[0], days[1]}, "months": {days[0][:7]}}):
        r = export_excel(con, site, out, **kw)
        assert r.other_site and (r.written, r.deleted, r.failed, r.unchanged, r.kept) == ([], [], [], 0, 0), kw
        assert not r.record_lost
        assert snapshot(out) == before
    text = format_result(r)
    assert "다른 사이트 팩의 사본" in text and OTHER_SITE not in text and "synthetic" not in text
    s = null_run.settings
    assert main(["export", "excel", str(out), "--site", str(s.site), "--work-root", str(s.work_root)]) == 1
    cap = capsys.readouterr()
    assert "다른 사이트 팩의 사본" in cap.out and len(cap.out.strip().splitlines()) == 1
    assert all(name not in cap.out + cap.err for name in (OTHER_SITE, "synthetic"))
    assert snapshot(out) == before


def test_an_old_record_without_a_site_is_accepted_and_gets_one(null_run, tmp_path):
    """옛 기록(사이트 없음)은 받아들인다 — 여느 내보내기이고(바뀐 것이 없으면 쓰지 않는다), 끝나면 기록에 이 사이트의 이름이 있다."""
    import json

    con, site = null_run.con, null_run.site
    out = tmp_path / "out"
    out.mkdir()
    export_excel(con, site, out, full=True)
    data = record_of(out)
    del data["site"]
    (out / RECORD_NAME).write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    xlsx = {k: v for k, v in snapshot(out).items() if k != RECORD_NAME}
    r = export_excel(con, site, out, full=True)
    assert not r.other_site and not r.record_lost and not r.failed
    assert r.written == [] and r.unchanged == len(all_files(con))
    assert {k: v for k, v in snapshot(out).items() if k != RECORD_NAME} == xlsx       # 엑셀 파일은 그대로
    again = record_of(out)
    assert again["site"] == "synthetic" and again["files"] == data["files"]


def test_a_site_pack_without_a_site_name_is_refused(site, null_run, tmp_path, capsys):
    """site.toml 에 [site] name 이 없으면: export_excel 은 ExportError, 명령은 한 줄로 거절, 자동 내보내기는 켜지 않는다 (한 줄 알림).
    폴더 이름(site)으로 대신하지 않는다 (4.1 다)."""
    import shutil

    from minedocscan.cli import _round_jobs, main
    from minedocscan.export.auto import AutoExport
    from minedocscan.touched import Touched

    root = tmp_path / "site"
    shutil.copytree(site.root, root)
    toml = root / "site.toml"
    text = toml.read_text(encoding="utf-8")
    assert text.count('name = "synthetic"\n') == 1
    toml.write_text(text.replace('name = "synthetic"\n', ""), encoding="utf-8")
    noname = SitePack(root)
    assert noname.declared_name is None and noname.name == "site"                     # 폴더 이름은 표식이 아니다
    out = tmp_path / "out"
    out.mkdir()
    with pytest.raises(ExportError) as e:
        export_excel(null_run.con, noname, out, full=True)
    assert "[site] name" in str(e.value)
    s = null_run.settings
    with pytest.raises(SystemExit) as e:
        main(["export", "excel", str(out), "--site", str(root), "--work-root", str(s.work_root)])
    msg = str(e.value)
    assert "[site] name" in msg and "\n" not in msg
    st = Settings(site=root, excel_dir=out, reviews=tmp_path / "reviews.jsonl")
    x = AutoExport(st, noname)
    assert not x.enabled and x.reason == "no_site_name" and x.status["reason"] == "no_site_name"
    assert "[site] name" in x.notice and "\n" not in x.notice
    assert x.after_round(null_run.con, Touched(everything=True)) is None
    capsys.readouterr()
    assert _round_jobs(st, noname)["excel"].reason == "no_site_name"
    err = capsys.readouterr().err
    assert "[site] name" in err and len([ln for ln in err.splitlines() if "엑셀" in ln]) == 1
    assert list(out.iterdir()) == []


def test_an_empty_work_db_does_not_clean_the_folder(null_run, tmp_path):
    """쪽이 하나도 없는 작업 DB 로 전부 훑어도 기록된 파일을 지우지 않는다 — "남겨 둔 파일"로 센다 (4.1 다)."""
    from minedocscan.intake.worker import format_round
    from minedocscan.store.db import open_db

    con, site = null_run.con, null_run.site
    out = tmp_path / "out"
    out.mkdir()
    export_excel(con, site, out, full=True)
    before = snapshot(out)
    n = len(list(out.rglob("*.xlsx")))
    assert n == len(all_files(con))
    empty = open_db(f"sqlite:///{tmp_path / 'empty' / 'minedocscan.db'}")
    try:
        assert empty.execute("SELECT COUNT(*) FROM doc_page").fetchone()[0] == 0
        r = export_excel(empty, site, out, full=True)
    finally:
        empty.close()
    assert r.empty_db and not r.deleted and not r.written and not r.failed and r.kept == n
    assert snapshot(out) == before                                       # 기록 파일도 그대로
    assert "작업 DB 에 쪽이 없어" in format_result(r)
    assert "작업 DB 에 쪽이 없어" in format_round({"excel": r.as_dict()})
    r = export_excel(con, site, out, full=True)                          # 쪽이 있는 DB 로 돌아오면 여느 때처럼
    assert not r.empty_db and r.written == [] and r.unchanged == n


def test_date_selection_is_honoured_on_a_folder_with_a_record(world, tmp_path, capsys):
    """기록 파일이 있는 폴더에서 --date·--month·--from/--to: 고른 날짜(달)의 파일만 다시 쓴다 — 전부 훑지 않는다 (4.1 자).
    두 달에 걸치게 d(01-08)를 2030-02-08 로 옮기고, 고르지 않은 쪽에도 늘 바뀐 칸을 남겨 둔다 (전부 훑으면 그 파일도 쓴다)."""
    import json

    from minedocscan.cli import main
    from minedocscan.intake import decisions as decs

    pipe, site, st = world["pipe"], world["site"], world["st"]
    con = pipe.con
    decs.save(con, st.decisions_path(site.root), [{"target": world["ids"]["d_2030-01-08"], "kind": "date",
                                                   "value": "2030-02-08"}], "jp")
    assert pipe.process_pending() == 1
    x, y = "2030-01-07", "2030-02-08"
    assert y in days_of(con) and x in days_of(con)
    out = tmp_path / "엑셀"
    out.mkdir()
    common = ["--site", str(st.site), "--work-root", str(st.work_root), "--json"]
    used: set[str] = set()

    def export(*args) -> list[str]:
        assert main(["export", "excel", str(out), *common, *args]) == 0
        return sorted(json.loads(capsys.readouterr().out)["written"])

    def change(day: str) -> None:
        """그 날짜의 운반 횟수 칸 하나를 검수한다 — 그 날짜의 일별 파일과 그 달의 파일이 바뀐다."""
        fid = next(r[0] for r in con.execute("SELECT source_field_id FROM prod_haul WHERE work_date = ? AND source_role = 'log' "
                                             "AND review_status = 'pending' ORDER BY haul_id", (day,)) if r[0] not in used)
        used.add(fid)
        save(con, site, st, review_from_field(con, fid, "value", "3", "jp"))

    assert export() == all_files(con)                                   # 처음: 전부
    files_x, files_y = [daily_path(x), monthly_path(x[:7])], [daily_path(y), monthly_path(y[:7])]
    change(x)
    change(y)
    assert export("--date", x) == files_x                               # y 의 바뀐 칸은 남는다
    change(x)
    assert export("--month", y[:7]) == files_y                          # x 의 바뀐 칸은 남는다
    change(y)
    assert export("--from", "2030-01-01", "--to", "2030-01-31") == files_x
    assert export() == files_y                                          # 남겨 둔 y 의 바뀐 칸은 전부 훑기가 쓴다
    assert export() == []


def test_export_command_date_arguments(null_run, tmp_path, capsys):
    """--month 9999-12(date 가 다룰 수 있는 마지막 달)는 여느 실행이다 (OverflowError 가 아니다). --date 는 YYYY-MM-DD 만
    (20300107 은 거절), --month 는 YYYY-MM 만 (4.1 사)."""
    from minedocscan.cli import main

    s = null_run.settings
    out = tmp_path / "out"
    out.mkdir()
    common = ["export", "excel", str(out), "--site", str(s.site), "--work-root", str(s.work_root)]
    assert main(common) == 0
    capsys.readouterr()
    before = snapshot(out)
    for args in (["--month", "9999-12"], ["--date", "9999-12-31"], ["--from", "9999-12-01", "--to", "9999-12-31"]):
        assert main([*common, *args]) == 0, args
        assert "쓴 파일 0" in capsys.readouterr().out
    assert snapshot(out) == before
    for args, want in ((["--date", "20300107"], "YYYY-MM-DD"), (["--from", "20300107"], "YYYY-MM-DD"),
                       (["--month", "2030-1"], "YYYY-MM"), (["--month", "203001"], "YYYY-MM"),
                       (["--month", "2030-13"], "YYYY-MM")):
        with pytest.raises(SystemExit) as e:
            main([*common, *args])
        assert want in str(e.value), args
    assert snapshot(out) == before
